#!/usr/bin/env python3
"""Copy the photographs into images/photos and draw the history timeline.

Sources are given by environment variables, so that no private path is written into the
repository: PHOTOS is the root of the photo archive (with UiO/fourMs, UiO/RITMO/årsrapport,
UiO/RITMO/best and arj/), REPORT_IMAGES is the image folder of RITMO's final report, and
INTRO_PPTX is the earlier fourMs Lab introduction deck, whose pictures are taken by position.
Photographs wider than 1920 px are scaled down.

    PHOTOS=<photo archive> REPORT_IMAGES=<report>/images .venv/bin/python tools/make_figures.py
"""
import csv
import os
from pathlib import Path

import matplotlib
matplotlib.use("Agg")
import matplotlib.pyplot as plt
from PIL import Image, ImageOps

HERE = Path(__file__).resolve().parent.parent
IMG = HERE / "images"

RED, GREEN, YELLOW, INK = "#B22309", "#9FD54F", "#F2DF30", "#231310"
FONT = ["Nimbus Sans", "DejaVu Sans"]

PHOTOS = {
    # the lab
    "UiO/fourMs/3d9a8069.jpg": "concert-in-lab.jpg",
    "UiO/fourMs/3d9a8013.jpg": "quartet-mocap.jpg",
    "UiO/fourMs/3d9a8005.jpg": "control-room.jpg",
    "UiO/fourMs/20210528-P1022682_1000px.JPG": "cameras.jpg",
    "UiO/fourMs/20210528-P1022685_1000px.JPG": "camera-ring.jpg",
    "UiO/fourMs/20210528-P1022679_1000px.JPG": "markers.jpg",
    "UiO/fourMs/2019_04_01a.jpg": "mocap-suit.jpg",
    "UiO/fourMs/emg_sensor_placement.jpg": "emg.jpg",
    "UiO/fourMs/5r0a9097.jpg": "open-lab.jpg",
    "UiO/fourMs/2017-06-01-DSC_9528.JPG": "mocap-screen.jpg",
    "UiO/fourMs/20210502-P1022609.JPG": "preparing.jpg",
    "UiO/fourMs/20220524_p1026592.jpg": "analysis-screens.jpg",
    "UiO/fourMs/20230921_161957.jpg": "projection-dance.jpg",
    "UiO/fourMs/2025-03-07-MUS2721 Performance research.png": "piano-course.jpg",
    "UiO/fourMs/04-20190304-dsc_3799.jpg": "winter-school.jpg",
    "UiO/fourMs/20240606_144711.jpg": "kork-recording.jpg",
    "UiO/fourMs/20240403_172119.jpg": "experiment-group.jpg",
    # history
    "UiO/fourMs/intermedia.png": "history-2005.png",
    "UiO/fourMs/24-channel.jpg": "history-2008.jpg",
    "UiO/fourMs/2012-06-12-dscn2710.jpg": "history-2011.jpg",
    # MusicLab and standstill
    "UiO/RITMO/årsrapport/2024/Abels-KORK/11-586a1773-kopi_1920.jpg": "kork-violinist.jpg",
    "UiO/RITMO/årsrapport/2023/lydo/01-20230216_101518.jpg": "concert-stavanger.jpg",
    "arj/sverm/2012-05-10-sverm-standstill.JPG": "sverm-standstill.jpg",
    "UiO/RITMO/best/motion-history-images8.png": "motion-history.png",
}
REPORT = {"2025/aim-workshop-standstill-mocap.jpg": "standstill-lab.jpg"}
# From the earlier fourMs Lab introduction deck (INTRO_PPTX): (slide, picture on the slide) -> name.
INTRO = {
    (6, 1): "piano-cameras.jpg", (7, 1): "violin-voice.jpg",
    (11, 1): "standstill-championship.jpg", (12, 3): "standstill-championship-2.jpg",
    (13, 1): "standstill-database.png", (14, 1): "rowing.jpg", (15, 1): "guitars.jpg",
    (17, 2): "musiclab-workshop.jpg", (17, 3): "musiclab-concert.jpg", (17, 4): "musiclab-panel.jpg",
    (17, 5): "musiclab-data-jockey.jpg", (23, 1): "musiclab-copenhagen.jpg", (24, 1): "mooc-motion-capture.jpg",
    (28, 1): "dancing-in-lab.gif",
    (4, 1): "handbook-e-infrastructure.jpg", (4, 2): "handbook-infrastructure.jpg",
    (4, 3): "handbook-software.png", (4, 4): "handbook-database.jpg",
}


def copy(src, name):
    im = ImageOps.exif_transpose(Image.open(src))
    if im.width > 1920:
        im = im.resize((1920, round(im.height * 1920 / im.width)), Image.LANCZOS)
    dst = IMG / "photos" / name
    if dst.suffix == ".png":
        im.save(dst)
    else:
        im.convert("RGB").save(dst, quality=88)


def intro(pptx):
    import io
    from pptx import Presentation
    deck = Presentation(pptx)
    for i, slide in enumerate(deck.slides, 1):
        pics = [sh for sh in slide.shapes if sh.shape_type == 13]
        for k, sh in enumerate(pics, 1):
            name = INTRO.get((i, k))
            if not name:
                continue
            if name.endswith(".gif"):
                (IMG / "photos" / name).write_bytes(sh.image.blob)
            else:
                copy(io.BytesIO(sh.image.blob), name)


def timeline():
    """The lab's history on a true time axis. Labels alternate above and below the axis, and
    every other pair sits further out, so that events a year apart do not collide."""
    with open(HERE / "data/history.tsv", encoding="utf-8") as f:
        rows = list(csv.DictReader(f, delimiter="\t"))
    plt.rcParams["font.family"] = FONT
    fig, ax = plt.subplots(figsize=(19.2, 9.6), dpi=100)
    years = [int(r["year"]) for r in rows]
    x0, x1 = min(years) - 0.8, max(years) + 0.8
    ax.plot([x0, x1], [0, 0], color=INK, lw=3, solid_capstyle="round", zorder=1)
    for y in range(min(years), max(years) + 1):
        ax.plot([y, y], [-0.07, 0.07], color=INK, lw=1, zorder=1)
    colours = [GREEN, RED, YELLOW]
    for i, r in enumerate(rows):
        x = int(r["year"])
        up = 1 if i % 2 == 0 else -1
        h = 0.75 if (i // 2) % 2 == 0 else 2.35
        ax.plot([x, x], [0, up * h], color=INK, lw=1.4, zorder=1)
        ax.scatter([x], [0], s=420, color=colours[i % 3], edgecolor=INK, linewidth=2, zorder=2)
        ax.text(x, up * (h + 0.05), r["year"], ha="center", va="bottom" if up > 0 else "top",
                fontsize=28, fontweight="bold", color=INK)
        words, lines, line = r["event"].split(), [], ""
        for w in words:
            if line and len(line) + len(w) > 17:
                lines.append(line)
                line = w
            else:
                line = (line + " " + w).strip()
        lines.append(line)
        ax.text(x, up * (h + (0.62 if up > 0 else 0.48)), "\n".join(lines), ha="center",
                va="bottom" if up > 0 else "top", fontsize=15, color=INK, linespacing=1.2)
    ax.set_xlim(x0 - 0.6, x1 + 0.6)
    ax.set_ylim(-4.3, 4.3)
    ax.axis("off")
    fig.savefig(IMG / "timeline.png", bbox_inches="tight", pad_inches=0.15, facecolor="white")


def main():
    (IMG / "photos").mkdir(parents=True, exist_ok=True)
    root = os.environ.get("PHOTOS")
    if root:
        for src, name in PHOTOS.items():
            copy(Path(root) / src, name)
    rep = os.environ.get("REPORT_IMAGES")
    if rep:
        for src, name in REPORT.items():
            copy(Path(rep) / src, name)
    pptx = os.environ.get("INTRO_PPTX")
    if pptx:
        intro(pptx)
    timeline()
    print("photos:", len(list((IMG / "photos").iterdir())))


if __name__ == "__main__":
    main()
